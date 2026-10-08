"""Create an illustrated academic narrative from Gia Han's Spotify notebook."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "reports" / "chart_assets"
OUTPUT = ROOT / "reports" / "Bao_cao_khai_pha_insight_Spotify_Gia_Han_co_hinh.docx"


def page_field(paragraph):
    run = paragraph.add_run()
    for kind, value in (("begin", None), (None, "PAGE"), ("separate", None), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = value
        run._r.append(element)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(.79)
sec.top_margin = Inches(.70)
sec.bottom_margin = Inches(.68)
sec.header_distance = Inches(.30)
sec.footer_distance = Inches(.31)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(20, 20, 20)
normal.paragraph_format.line_spacing = 1.14
normal.paragraph_format.space_after = Pt(6)
for name, size, before, after in (
    ("Title", 18.5, 0, 10),
    ("Heading 1", 14, 17, 7),
    ("Heading 2", 11.5, 12, 6),
):
    style = doc.styles[name]
    style.font.name = "Aptos Display" if name != "Heading 2" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

subtitle = doc.styles["Subtitle"]
subtitle.font.name = "Aptos"
subtitle.font.size = Pt(10.5)
subtitle.font.italic = False
subtitle.font.color.rgb = RGBColor(0, 0, 0)

# The default Word Title style sometimes carries a blue rule.
title_style_ppr = doc.styles["Title"].element.get_or_add_pPr()
style_border = title_style_ppr.find(qn("w:pBdr"))
if style_border is not None:
    title_style_ppr.remove(style_border)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("GIA HÂN  |  PHÂN TÍCH SPOTIFY CHARTS")
r.font.name = "Aptos"
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(0, 0, 0)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("Báo cáo khai phá insight  ·  ")
r.font.name = "Aptos"
r.font.size = Pt(8)
page_field(footer)


def para(text, *, keep=False):
    p = doc.add_paragraph(text)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = keep
    return p


def figure(number, heading, caption, paragraphs, width=6.6, display_note=None):
    doc.add_heading(heading, 2)
    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_after = Pt(2)
    image_p.paragraph_format.keep_with_next = True
    picture = image_p.add_run().add_picture(str(ASSETS / f"figure_{number:02}.png"), width=Inches(width))
    picture._inline.docPr.set("descr", f"Biểu đồ {number}: {caption}")

    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    cp.paragraph_format.keep_with_next = True
    cr = cp.add_run(f"Hình {number}. {caption}")
    cr.italic = True
    cr.font.size = Pt(8.5)
    cr.font.color.rgb = RGBColor(55, 55, 55)
    if display_note:
        nr = cp.add_run(f" {display_note}")
        nr.italic = True
        nr.font.size = Pt(8.5)
        nr.font.color.rgb = RGBColor(55, 55, 55)

    for i, text in enumerate(paragraphs):
        para(text, keep=(i == 0 and len(paragraphs) > 1))


doc.add_heading("Phân tích đặc trưng âm thanh và xu hướng Spotify Charts", 0)
p = doc.add_paragraph("Báo cáo khai phá insight từ ba dashboard và chín biểu đồ")
p.style = "Subtitle"
para("Người thực hiện: Gia Hân")

doc.add_heading("Tóm tắt", 1)
para(
    "Báo cáo khảo sát hồ sơ âm thanh của các ca khúc xuất hiện trên Spotify Charts, sự khác biệt giữa thị trường quốc gia và biến động theo thời gian. "
    "Dữ liệu cho thấy Top 10 có xu hướng nhanh và dễ nhảy hơn các nhóm hạng thấp, song mức chênh lệch theo từng đặc trưng tương đối nhỏ. "
    "Bản đồ quốc gia và biểu đồ theo tháng lại cho thấy bối cảnh địa lý, mùa vụ có thể làm thay đổi đáng kể giá trị trung bình của các chỉ số. "
    "Vì vậy, các kết quả nên được đọc như những liên hệ quan sát, chưa phải công thức tạo hit hay bằng chứng về tác động nhân quả."
)

doc.add_heading("Dữ liệu và nguyên tắc diễn giải", 1)
para(
    "Chín hình được trích xuất trực tiếp từ cấu hình Plotly lưu trong notebooks/04_story.ipynb. "
    "Các số liệu trong phần nhận xét được đối chiếu với bảng dữ liệu đã xử lý: tracks_dimension, charts_clean, mart_artist_chart, "
    "mart_country_year, mart_audio_trend_weekly và covid_weekly_clean. Phạm vi thời gian và đơn vị phân tích không giống nhau giữa các hình: "
    "có hình ở cấp lượt hiện diện trên bảng xếp hạng, có hình ở cấp ca khúc, nghệ sĩ, quốc gia–năm hoặc tuần. "
    "Trung bình có trọng số stream phản ánh cơ cấu lượt nghe trong dữ liệu BXH; trung bình không trọng số phản ánh cơ cấu các mục quan sát."
)
para(
    "Các hình có menu tương tác được trình bày dưới dạng ảnh tĩnh ở trạng thái mặc định: Energy ở Hình 2, Valence ở Hình 6 và Acousticness ở Hình 9. "
    "Các nhận xét về chỉ số khác của cùng hình dựa trên dữ liệu gốc, không phải trên ảnh tĩnh duy nhất. "
    "Trong toàn báo cáo, phần mô tả số liệu được tách khỏi cơ chế giả thuyết; những biến chưa được đo, như phong tỏa, thuật toán đề xuất hoặc động cơ cảm xúc của người nghe, không được xem là nguyên nhân đã chứng minh."
)

doc.add_heading("Dashboard 1 Công thức tạo hit và xu hướng âm nhạc thời đại", 1)
para(
    "Dashboard thứ nhất xem xét khác biệt âm thanh giữa các tầng thứ hạng, phân phối theo nhãn thể loại, mức tập trung stream ở nghệ sĩ, "
    "hình thức hợp tác và thời lượng ca khúc. Tên gọi “công thức tạo hit” mang tính định hướng câu hỏi; năm biểu đồ không tạo thành mô hình dự báo.",
    keep=True,
)

figure(
    1,
    "Hồ sơ âm thanh theo tầng thứ hạng",
    "Chênh lệch chuẩn hóa của chín đặc trưng âm thanh giữa bốn tầng hạng Global giai đoạn 2020–2022.",
    [
        "Sau khi chuẩn hóa từng đặc trưng theo trung bình và độ lệch chuẩn của toàn bộ lượt hiện diện trên BXH Global giai đoạn 2020–2022, "
        "nhóm Top 1–10 nổi bật ở tempo (+0,16 độ lệch chuẩn), liveness (+0,14), danceability (+0,13), valence (+0,12) và energy (+0,11). "
        "Acousticness, speechiness và instrumentalness của nhóm này thấp hơn mức chung khoảng 0,11–0,12 độ lệch chuẩn; loudness gần như không đổi. "
        "So với hạng 101–200, Top 10 nhanh hơn 5,53 BPM, cao hơn 0,026 điểm danceability và 0,035 điểm valence. Những chênh lệch này nhất quán về hướng, nhưng quy mô vừa phải.",
        "Hồ sơ trên phù hợp với khả năng các bài hát có nhịp rõ và cấu trúc dễ vận động đáp ứng nhiều bối cảnh nghe phổ biến. "
        "Tuy nhiên, dữ liệu không đo ngữ cảnh nghe, playlist hoặc chiến lược sản xuất, nên không thể quy khác biệt thứ hạng cho các cơ chế đó. "
        "Hình này là ma trận mức lệch z-score, không phải ma trận tương quan giữa các đặc trưng. Mỗi ca khúc có thể lặp lại qua nhiều tuần; "
        "do đó, kết quả mô tả hồ sơ của lượt hiện diện trên BXH hơn là tác động độc lập của một thuộc tính lên khả năng thành hit."
    ],
)

figure(
    2,
    "Phân phối đặc trưng theo nhãn thể loại",
    "Boxplot của năm nhãn thể loại được chọn; ảnh tĩnh hiển thị Energy.",
    [
        "Trong 2.908 ca khúc thuộc năm nhãn thể loại được chọn, K-Pop và Latin có trung vị energy cao, lần lượt là 0,732 và 0,715, "
        "trong khi Pop đạt 0,604. Khi chuyển chỉ số trong bản tương tác, Latin và Rap dẫn đầu về danceability với trung vị 0,755 và 0,745. "
        "Latin cũng có valence trung vị 0,674, cao hơn Pop (0,446) và Rap (0,436); ngược lại, Rap có tempo trung vị 126,8 BPM, cao nhất trong năm nhóm. "
        "Các hộp phân phối vẫn chồng lấn, cho thấy nhãn thể loại không xác định duy nhất đặc trưng của từng bài hát.",
        "Sự phân hóa này phù hợp với khác biệt trong quy ước tiết tấu và phối khí giữa các cộng đồng thể loại, nhưng biểu đồ không đo văn hóa lễ hội, mục đích biểu diễn hay nội dung lời ca. "
        "Vì trường artist_genre là nhãn của nghệ sĩ, không phải phân loại riêng của từng ca khúc, diễn giải nên dừng ở mức so sánh nhóm được gắn nhãn. "
        "Không có thể loại nào cao trên cả bốn chiều; đánh giá bài hát theo nhiều đặc trưng cho kết quả chính xác hơn một nhãn “sôi động” chung."
    ],
    display_note="Ba chỉ số còn lại được phân tích từ dữ liệu của biểu đồ tương tác.",
)

figure(
    3,
    "Mức tập trung stream ở nghệ sĩ và thể loại",
    "Treemap của 80 nghệ sĩ dẫn đầu sau khi lọc bảy nhãn thể loại; G trên các ô tương ứng tỷ lượt stream.",
    [
        "Trong 80 nghệ sĩ được chọn, nhánh Pop gồm 36 nghệ sĩ và chiếm 46,1% tổng stream của mẫu. "
        "Hai nhánh Latin và trap latino cộng lại đạt khoảng 94,2 tỷ stream, tương ứng 26,6%; Bad Bunny có ô đơn lẻ lớn nhất với khoảng 19,9 tỷ stream. "
        "Sự khác biệt về diện tích ô cho thấy phân phối lượt nghe trong nhóm nghệ sĩ dẫn đầu không đồng đều và tập trung ở một số tên tuổi lớn.",
        "Quy mô danh mục bài hát, thời gian hiện diện trên BXH và khả năng tiếp cận nhiều thị trường đều có thể góp phần tạo nên chênh lệch này. "
        "Không thể từ treemap suy ra thuật toán đề xuất đã khuếch đại Pop hoặc Latin, vì hình không có biến đo mức hiển thị do thuật toán. "
        "Các tỷ trọng cũng chỉ có nghĩa trong mẫu 80 nghệ sĩ của bảy nhãn được lọc; gọi đó là thị phần của toàn bộ Spotify sẽ vượt quá phạm vi dữ liệu."
    ],
)

figure(
    4,
    "Hợp tác biểu diễn và quy mô stream",
    "Tỷ lệ ca khúc đơn ca và hợp tác cùng stream BXH trung bình trên mỗi bài.",
    [
        "Trong 39.190 ca khúc của tập dữ liệu, đơn ca chiếm 60,8% và hợp tác biểu diễn chiếm 39,2%. "
        "Dù ít hơn về số lượng, nhóm hợp tác có trung bình 53,7 triệu stream BXH mỗi bài, so với 17,7 triệu của nhóm đơn ca, tức xấp xỉ 3,0 lần. "
        "Trung vị của hai nhóm cũng khác nhau, khoảng 2,30 triệu và 0,70 triệu, cho thấy chênh lệch không chỉ do một vài siêu hit kéo giá trị trung bình lên.",
        "Việc kết hợp nhiều nghệ sĩ có thể mở rộng cộng đồng người hâm mộ và tạo thêm cơ hội tiếp cận ban đầu, nhưng hình không đo trực tiếp các kênh lan truyền này. "
        "Độ nổi tiếng sẵn có của nghệ sĩ, số tuần hiện diện, số quốc gia, thể loại và thời điểm phát hành đều có thể đồng thời ảnh hưởng đến stream. "
        "Vì vậy, kết quả mô tả mối liên hệ giữa hình thức hợp tác và quy mô lượt nghe trong nhóm bài đã lọt BXH, không chứng minh hợp tác làm tăng stream hay Spotify ưu tiên các ca khúc đó."
    ],
)

figure(
    5,
    "Thời lượng ca khúc trên bảng xếp hạng",
    "Thời lượng trung bình có trọng số stream theo tuần ở BXH Global giai đoạn 2017–2022.",
    [
        "Thời lượng trung bình có trọng số stream giảm từ 218,0 giây năm 2017 xuống 209,2 giây năm 2018, 201,2 giây năm 2019 và 198,5 giây năm 2020. "
        "Mức giảm 2017–2020 xấp xỉ 19,5 giây. Sau đó, chỉ số tăng trở lại lên 200,1 giây năm 2021 và 205,7 giây trong phần năm 2022 có dữ liệu. "
        "Đường trung bình trượt bốn tuần làm rõ giai đoạn rút ngắn, đồng thời cho thấy hiện tượng không tiếp tục theo một chiều đến cuối chuỗi.",
        "Thói quen tiêu thụ nội dung ngắn hoặc cấu trúc trả phí theo lượt nghe là những giả thuyết có thể đặt ra để giải thích giai đoạn giảm, "
        "nhưng bảng dữ liệu không chứa tỷ lệ bỏ qua bài, mức sử dụng TikTok hoặc quyết định rút ngắn bản phối. "
        "Do đó, không thể kết luận “TikTok-hóa” là nguyên nhân hay bài dưới ba phút đã trở thành chuẩn bắt buộc; mức trung bình quan sát vẫn gần 200 giây. "
        "Năm 2022 chưa đủ cả năm, nên so sánh cuối chuỗi cần được đọc thận trọng."
    ],
)

doc.add_heading("Dashboard 2 Bản sắc địa lý và hồ sơ âm thanh", 1)
para(
    "Dashboard thứ hai chuyển đơn vị phân tích từ ca khúc và nghệ sĩ sang quốc gia trong cùng năm 2021. "
    "Bản đồ mô tả trung bình có trọng số stream, còn radar Việt Nam–Nhật Bản mô tả trung bình không trọng số; "
    "hai hình bổ sung góc nhìn nhưng không được trộn trị số như cùng một phép tính.",
    keep=True,
)

figure(
    6,
    "Phân hóa địa lý của bốn đặc trưng âm thanh",
    "Bản đồ 68 quốc gia có dữ liệu Spotify Charts năm 2021; ảnh tĩnh hiển thị Valence.",
    [
        "Bản đồ ghi nhận sự phân hóa giữa các thị trường có dữ liệu. Brazil đạt valence có trọng số stream 0,690; Chile có danceability 0,752. "
        "Ở chiều acousticness, Việt Nam đạt 0,465, cao hơn đáng kể Nhật Bản (0,173), trong khi energy của Nhật Bản là 0,727, cao hơn Việt Nam (0,532). "
        "Trung bình valence của các nước Mỹ Latinh và Caribe trong mẫu khoảng 0,627; vì thế không nên khái quát rằng toàn bộ khu vực đều vượt ngưỡng 0,65.",
        "Cơ cấu ca khúc phổ biến, tỷ trọng nhạc nội địa và dòng nhạc quốc tế có thể cùng định hình hồ sơ âm thanh của từng quốc gia. "
        "Những khác biệt trên phù hợp với giả thuyết về thị hiếu thị trường, nhưng dữ liệu quốc gia tổng hợp không cho biết đặc điểm của từng người nghe và không xác nhận nguyên nhân văn hóa. "
        "Bản đồ chỉ bao phủ 68 quốc gia có mã địa lý trong năm 2021; màu trắng không nên được hiểu là quốc gia không nghe nhạc Spotify."
    ],
    display_note="Danceability, Acousticness và Energy được đọc từ ba trạng thái khác của bản tương tác.",
)

figure(
    7,
    "Đối chiếu Việt Nam và Nhật Bản",
    "Radar sáu đặc trưng theo trung bình không trọng số của hai BXH quốc gia năm 2021.",
    [
        "Radar cho thấy Việt Nam có acousticness trung bình 0,474, cao hơn Nhật Bản (0,172) gần 2,8 lần, trong khi energy là 0,535 so với 0,727. "
        "Valence của Việt Nam đạt 0,460, thấp hơn Nhật Bản (0,587), còn danceability nhỉnh hơn (0,663 so với 0,606). "
        "Sự khác biệt theo từng trục cho thấy ngay trong nhóm thị trường châu Á, hồ sơ ca khúc thịnh hành không đồng nhất.",
        "Kết quả phù hợp với khả năng hai thị trường có cơ cấu ca khúc và phong cách sản xuất phổ biến khác nhau. "
        "Tuy vậy, acousticness chỉ là đặc trưng âm thanh được ước lượng, không chứng minh ca khúc dùng guitar hoặc piano; "
        "radar cũng không xác nhận Pop Ballad hay Anime OST là nguyên nhân. "
        "Hình này dùng trung bình không trọng số, nên không so sánh trực tiếp 0,474 với trị số 0,465 của Việt Nam trên bản đồ có trọng số stream."
    ],
    width=5.3,
)

doc.add_heading("Dashboard 3 Mùa vụ và giai đoạn COVID 19", 1)
para(
    "Dashboard thứ ba khảo sát hai dạng biến thiên thời gian khác nhau: chu kỳ tháng trong nhiều năm và sự cùng biến động theo tuần ở giai đoạn COVID-19. "
    "Cả hai phép so sánh đều có nguy cơ bị chi phối bởi thành phần ca khúc, xu hướng dài hạn và mùa vụ.",
    keep=True,
)

figure(
    8,
    "Biến động theo tháng của danceability và acousticness",
    "Trung bình có trọng số stream theo tháng lịch của BXH Global trong các tuần có dữ liệu.",
    [
        "Danceability duy trì quanh 0,69–0,70 trong các tháng 1–11 rồi giảm còn 0,661 ở tháng 12. "
        "Cùng thời điểm, acousticness tăng từ mức thông thường khoảng 0,22–0,24 lên 0,294. "
        "Sự dịch chuyển ngược chiều của hai chỉ số khiến tháng 12 trở thành điểm lệch mùa vụ rõ nhất trong biểu đồ gộp.",
        "Sự quay lại của danh mục nhạc dịp cuối năm là cơ chế có thể giải thích đồng thời cả hai hướng biến động; "
        "tuy nhiên, hình chưa định danh bài Giáng sinh hoặc so sánh tỷ trọng ca khúc cũ giữa các tháng. "
        "Do các tháng được gộp qua những năm có mức độ bao phủ không giống nhau, kết luận về tính lặp lại hàng năm cần được kiểm tra bằng đường theo từng năm, "
        "sau đó tách ảnh hưởng của thể loại và thị trường."
    ],
)

figure(
    9,
    "Âm thanh và số ca tử vong COVID 19 theo tuần",
    "Số ca tử vong Global và Acousticness có trọng số stream trong 130 tuần ghép khớp, 2020–2022.",
    [
        "Số ca tử vong tuần đạt đỉnh 102.385 vào ngày 28/01/2021; cùng tuần, acousticness của BXH Global là 0,283, energy 0,607 và valence 0,513. "
        "Trong toàn kỳ, acousticness dao động 0,214–0,367 và energy 0,576–0,668. Valence có trung bình 0,523 nhưng từng xuống 0,480, "
        "nên nhận định chỉ số này luôn trên 0,50 là không chính xác. Hệ số tương quan Pearson cùng tuần với số tử vong lần lượt khoảng 0,345 cho acousticness, −0,164 cho energy và 0,288 cho valence; "
        "đây là liên hệ yếu đến vừa, chưa điều chỉnh yếu tố gây nhiễu.",
        "Sáu đặc trưng còn lại không cho thấy một phản ứng đồng nhất: danceability nằm trong dải 0,625–0,724; loudness từ −7,50 đến −5,71 dB; "
        "speechiness 0,069–0,133; instrumentalness 0,004–0,024; tempo 117,2–125,6 BPM; và liveness 0,151–0,201. "
        "Mức biến thiên hẹp của một số chỉ số gợi ý cần tách từng đặc trưng, thay vì gộp thành khẳng định rằng toàn bộ âm nhạc đã thay đổi mạnh trong dịch.",
        "Hồ sơ “êm hơn nhưng không nhất thiết buồn hơn” phù hợp với một số giai đoạn, song biểu đồ hai trục không kiểm định được việc tử vong hoặc phong tỏa gây thay đổi lựa chọn âm nhạc. "
        "Dữ liệu không có biến về chính sách hạn chế, hoàn cảnh nghe hay trạng thái tâm lý. "
        "Các phân tích tiếp theo nên kiểm soát mùa vụ và xu hướng năm, thử độ trễ một đến hai tuần và lặp lại ở cấp quốc gia trước khi bàn về cơ chế thích nghi."
    ],
    display_note="Tám đặc trưng khác được phân tích từ các trạng thái lựa chọn của bản tương tác.",
)

doc.add_heading("Kết luận và hướng kiểm định tiếp theo", 1)
para(
    "Ba dashboard tạo ra một tuyến phân tích từ đặc điểm của bài hát dẫn đầu đến sự khác biệt thị trường và biến động theo thời gian. "
    "Nhóm Top 10 có hồ sơ hơi thiên về nhịp nhanh, dễ nhảy; nghệ sĩ lớn và các bản hợp tác gắn với lượng stream cao hơn; "
    "quốc gia và tháng lịch lại làm thay đổi đáng kể giá trị trung bình. "
    "Những kết quả này giúp đặt giả thuyết, nhưng chưa xác định cấu trúc nhân quả của một bản hit."
)
para(
    "Để nghiên cứu khả năng dự đoán, cần chuyển sang dữ liệu ca khúc–tuần, xác định trước biến đích như vào Top 10 ở tuần kế tiếp, "
    "chia tập huấn luyện và kiểm tra theo thời gian, đồng thời kiểm soát nghệ sĩ, thị trường, mùa vụ, thời gian hiện diện và mức stream trước đó. "
    "Đối với các giả thuyết về TikTok, Giáng sinh hoặc phong tỏa, nên bổ sung biến đo trực tiếp thay vì dùng diễn giải từ xu hướng đồ thị."
)

doc.core_properties.title = "Phân tích đặc trưng âm thanh và xu hướng Spotify Charts"
doc.core_properties.author = "Gia Hân"
doc.core_properties.subject = "Báo cáo khoa học minh họa ba dashboard Spotify"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
