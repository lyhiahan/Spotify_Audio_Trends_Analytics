"""Build the scientific insight report from verified notebook findings."""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(__file__).with_name("Bao_cao_khai_pha_insight_Spotify_Gia_Han.docx")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, text, sep, end):
        run._r.append(node)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.72)
sec.bottom_margin = Inches(.67)
sec.left_margin = Inches(.85)
sec.right_margin = Inches(.85)
sec.header_distance = Inches(.32)
sec.footer_distance = Inches(.33)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(25, 25, 25)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.10
for name, size, before, after in [
    ("Title", 19, 0, 8),
    ("Heading 1", 14, 12, 6),
    ("Heading 2", 11.5, 9, 3),
]:
    s = styles[name]
    s.font.name = "Aptos Display" if name != "Heading 2" else "Aptos"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
subtitle = styles["Subtitle"]
subtitle.font.name = "Aptos"
subtitle.font.size = Pt(10.5)
subtitle.font.italic = False
subtitle.font.color.rgb = RGBColor(65, 65, 65)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("GIA HÂN  |  SPOTIFY AUDIO TRENDS")
r.font.name = "Aptos"
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(80, 80, 80)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("Báo cáo khai phá insight  •  ")
r.font.size = Pt(8)
add_field(footer, "PAGE")

title_p = doc.add_heading("Khai phá insight từ ba dashboard Spotify", 0)
title_p.style = "Title"
title_p_pr = title_p._p.get_or_add_pPr()
border = title_p_pr.find(qn("w:pBdr"))
if border is not None:
    title_p_pr.remove(border)
style_border = styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr"))
if style_border is not None:
    styles["Title"].element.get_or_add_pPr().remove(style_border)
p = doc.add_paragraph()
p.add_run("Nghiên cứu và trực quan hóa đặc trưng âm thanh trên bảng xếp hạng Spotify").bold = True
p = doc.add_paragraph("Người thực hiện: Gia Hân  |  Bản phân tích đối chiếu với 04_story.ipynb và dữ liệu đã xử lý")
p.style = "Subtitle"

doc.add_heading("Tóm tắt", 1)
doc.add_paragraph(
    "Ba dashboard kể một chuỗi câu chuyện: hồ sơ âm thanh và cấu trúc thị trường của ca khúc lọt bảng xếp hạng; "
    "khác biệt giữa các quốc gia; và biến thiên theo mùa, theo giai đoạn COVID-19. "
    "Các kết quả nổi bật gồm: Top 10 nghiêng nhẹ về nhạc dễ nhảy và tempo nhanh, nhưng không khác biệt đáng kể về loudness; "
    "Việt Nam và Nhật Bản có hồ sơ âm thanh khác nhau rõ rệt; tháng 12 nổi bật với acousticness cao và danceability thấp. "
    "Những khác biệt này là liên hệ quan sát trong dữ liệu bảng xếp hạng, không chứng minh nguyên nhân hay dự đoán chắc chắn một bài hát sẽ thành hit."
)

doc.add_heading("Dữ liệu, đơn vị phân tích và cách đọc", 1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for c, text in zip(table.rows[0].cells, ["Dashboard", "Dữ liệu / phạm vi", "Đơn vị phân tích"]):
    c.text = text
    shade(c, "E8E8E8")
    for p in c.paragraphs:
        for run in p.runs:
            run.bold = True
for vals in [
    ("1 — Hit & xu hướng", "Global 2020–2022 cho heatmap; các bảng tracks, artist và weekly cho biểu đồ còn lại", "Lượt hiện diện trên BXH, ca khúc, nghệ sĩ hoặc tuần — tùy biểu đồ"),
    ("2 — Địa lý", "mart_country_year, năm 2021; 68 quốc gia có mã bản đồ", "Quốc gia–năm"),
    ("3 — Mùa vụ & COVID", "Global theo tuần; COVID ghép theo cùng mốc tuần, 130 tuần khớp", "Tuần hoặc tháng gộp từ các tuần có dữ liệu"),
]:
    row = table.add_row()
    for c, v in zip(row.cells, vals):
        c.text = v
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
doc.add_paragraph(
    "Các trị số *_weighted là trung bình có trọng số theo lượt stream; *_mean là trung bình không trọng số. "
    "Các chỉ số 0–1 mô tả thuộc tính âm thanh do Spotify gán, không phải đo trực tiếp cảm xúc hoặc động cơ của người nghe. "
    "Lượt stream trong báo cáo là lượt ghi nhận trong dữ liệu bảng xếp hạng, không đại diện cho toàn bộ lượt nghe trên nền tảng."
)


def insight(no, title, phenomenon, mechanism, conclusion, limit):
    doc.add_heading(f"Biểu đồ {no}. {title}", 2)
    for label, value in [
        ("Hiện tượng. ", phenomenon),
        ("Cơ chế khả dĩ. ", mechanism),
        ("Kết luận. ", conclusion),
        ("Giới hạn. ", limit),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        p.add_run(label).bold = True
        p.add_run(value)


doc.add_heading("Dashboard 1 — Công thức tạo hit & xu hướng âm nhạc thời đại", 1)
doc.add_paragraph(
    "Câu hỏi dẫn dắt: ca khúc ở nhóm thứ hạng cao khác gì về âm thanh, thể loại, mạng lưới nghệ sĩ và thời lượng? "
    "Ở đây, “công thức tạo hit” là cách đặt vấn đề khám phá, không phải công thức nhân quả hay mô hình dự báo."
)
insight(
    1, "Hồ sơ z-score theo bốn tầng thứ hạng",
    "Trong Global 2020–2022, Top 1–10 cao hơn trung bình BXH ở tempo (+0,16 độ lệch chuẩn), danceability (+0,13) "
    "và valence (+0,12), thấp hơn ở acousticness (−0,11); loudness gần như không khác biệt (−0,00). "
    "So với hạng 101–200, Top 10 nhanh hơn 5,53 BPM và cao hơn 0,026 điểm danceability.",
    "Nhịp nhanh và cấu trúc dễ vận động có thể phù hợp hơn với nhiều bối cảnh nghe đại chúng; đây là giả thuyết về mức độ phù hợp, chưa phải cơ chế được kiểm định.",
    "Nhóm đầu BXH có hồ sơ âm thanh hơi thiên về nhịp điệu, nhưng các mức lệch chỉ khoảng 0,1–0,2 độ lệch chuẩn; không có một đặc trưng đơn lẻ quyết định thứ hạng.",
    "Heatmap là ma trận mức chênh lệch chuẩn hóa, không phải ma trận tương quan. z = (trung bình nhóm thứ hạng − trung bình toàn BXH) / độ lệch chuẩn toàn BXH. "
    "Một ca khúc có thể xuất hiện nhiều tuần trên BXH, nên đây là hồ sơ của lượt hiện diện, không phải tập bài hát độc lập."
)
insight(
    2, "Phân phối bốn đặc trưng theo năm nhãn thể loại",
    "Trong 2.908 ca khúc thuộc năm nhãn được chọn, Latin có trung vị danceability 0,755 và valence 0,674; K-Pop cao nhất về energy (0,732); "
    "Rap có tempo trung vị 126,8 BPM. Pop có energy trung vị 0,604.",
    "Các khác biệt có thể liên quan đến quy ước sáng tác, sản xuất và đối tượng nghe của từng thể loại; boxplot còn cho thấy mức chồng lấn giữa các phân phối.",
    "Không có thể loại nào dẫn đầu ở mọi chiều; so sánh nhiều đặc trưng hữu ích hơn gắn một nhãn “sôi động” hay “tươi vui” chung cho cả thể loại.",
    "artist_genre là nhãn thể loại của nghệ sĩ, không nhất thiết là thể loại riêng của từng bài. Chưa có kiểm định ý nghĩa thống kê hoặc biến lời ca để suy ra nội dung tự sự."
)
insight(
    3, "Cơ cấu stream của 80 nghệ sĩ được chọn",
    "Trong 80 nghệ sĩ dẫn đầu sau khi lọc bảy nhãn thể loại, Pop chiếm 46,1% tổng stream của mẫu; Latin cộng trap latino chiếm 26,6% "
    "(94,2 tỷ). Bad Bunny là nghệ sĩ có ô lớn nhất, khoảng 19,9 tỷ stream trong dữ liệu.",
    "Sự tích tụ lượt nghe có thể phản ánh quy mô danh mục ca khúc, thời gian hiện diện và độ phủ đa thị trường của nghệ sĩ; biểu đồ không tách riêng tác động thuật toán gợi ý.",
    "Trong nhóm nghệ sĩ được chọn, lưu lượng tập trung vào một số thể loại và tên tuổi lớn; thị trường BXH có tính bất đối xứng rõ rệt.",
    "Đây là tỷ trọng trong mẫu 80 nghệ sĩ của bảy nhãn được lọc, không phải thị phần của toàn Spotify hoặc toàn bộ thể loại âm nhạc."
)
insight(
    4, "Đơn ca và hợp tác biểu diễn",
    "Trong 39.190 ca khúc của tập dữ liệu, đơn ca chiếm 60,8%; hợp tác chiếm 39,2%. "
    "Tổng stream BXH trung bình mỗi bài của nhóm hợp tác đạt 53,7 triệu, so với 17,7 triệu ở nhóm đơn ca (xấp xỉ 3,0 lần).",
    "Việc kết hợp nhiều nghệ sĩ có thể mở rộng tệp người hâm mộ và phạm vi tiếp cận ban đầu. Đây là lời giải thích khả dĩ, không phải kết quả đo lường kênh lan truyền.",
    "Hợp tác gắn với mức stream quan sát cao hơn trong tập bài đã vào BXH; điều này gợi ý một hướng phân tích chiến lược phát hành.",
    "So sánh chưa kiểm soát độ nổi tiếng, thời gian xuất hiện, số thị trường, năm phát hành hoặc thể loại. Không thể kết luận hợp tác làm tăng stream, "
    "cũng không thể khẳng định thuật toán Spotify ưu tiên bài hợp tác."
)
insight(
    5, "Xu hướng thời lượng bài hát, 2017–2022",
    "Thời lượng trung bình có trọng số stream của Global giảm từ 218,0 giây năm 2017 xuống 198,5 giây năm 2020 (−19,5 giây), "
    "sau đó tăng lên 200,1 giây năm 2021 và 205,7 giây năm 2022.",
    "Định dạng nghe trực tuyến và cạnh tranh sự chú ý có thể góp phần tạo động lực cho ca khúc cô đọng hơn; biểu đồ không có dữ liệu về TikTok, "
    "tỷ lệ nghe hết hoặc quyết định sản xuất để kiểm định giả thuyết này.",
    "Có một pha rút ngắn rõ trong 2017–2020, nhưng không phải xu hướng giảm liên tục đến 2022. Nên dùng tên trung tính “Xu hướng thời lượng ca khúc trên BXH”.",
    "Mức trung bình gần 200 giây vẫn trên ba phút; dữ liệu không cho phép nói ca khúc dưới ba phút đã trở thành quy tắc bắt buộc "
    "hoặc quy tắc tính stream 30 giây là nguyên nhân trực tiếp. Năm 2022 chưa đủ cả năm trong dữ liệu."
)

doc.add_heading("Dashboard 2 — Bản sắc địa lý và hồ sơ âm thanh", 1)
doc.add_paragraph("Câu hỏi dẫn dắt: khác biệt quốc gia có tồn tại ngay trong cùng một năm quan sát, và có thể mô tả cụ thể đến đâu?")
insight(
    6, "Bản đồ bốn đặc trưng theo quốc gia",
    "Bản đồ 2021 bao phủ 68 quốc gia có mã địa lý. Chile có danceability 0,752; Brazil có valence 0,690; "
    "Việt Nam có acousticness 0,465 và energy 0,532, còn Nhật Bản lần lượt là 0,173 và 0,727.",
    "Cơ cấu thể loại được nghe nhiều, thị trường nội địa và dòng ca khúc quốc tế có thể cùng tạo ra các hồ sơ quốc gia khác nhau. "
    "Các ví dụ trên phù hợp với giả thuyết về thị hiếu, nhưng chưa xác nhận nguyên nhân văn hóa.",
    "Bản đồ cho thấy độ phân hóa cụ thể giữa thị trường: Việt Nam thiên về âm thanh acoustic hơn Nhật Bản, trong khi Nhật Bản có energy cao hơn.",
    "Không phải mọi nước Mỹ Latinh đều đạt valence trên 0,65; trung bình khu vực trong mẫu khoảng 0,627. "
    "Dữ liệu chỉ gồm các nước có BXH Spotify và chỉ năm 2021; không đại diện toàn bộ người dân hay mọi thể loại nội địa."
)
insight(
    7, "Radar Việt Nam – Nhật Bản",
    "Với trung bình không trọng số năm 2021, Việt Nam có acousticness 0,474 và energy 0,535; Nhật Bản lần lượt 0,172 và 0,727. "
    "Valence của Việt Nam 0,460, thấp hơn Nhật Bản 0,587; danceability là 0,663 so với 0,606.",
    "Sự khác nhau có thể phản ánh tổ hợp ca khúc xuất hiện trên BXH ở hai thị trường, bao gồm cả nhạc trong nước và nhạc quốc tế.",
    "Ngay trong một tiểu vùng châu Á, hồ sơ nghe nhạc không đồng nhất; đối chiếu Việt Nam–Nhật Bản có sức giải thích hơn một nhận định chung về “gu châu Á”.",
    "Radar dùng *_mean không trọng số, còn bản đồ dùng *_weighted theo stream; không so sánh trực tiếp các con số của hai hình như cùng một thước đo. "
    "Acousticness là xác suất/đặc trưng ước lượng, không chứng minh ca khúc dùng nhạc cụ mộc; radar cũng không xác nhận thể loại Pop Ballad hay Anime OST là nguyên nhân."
)

doc.add_heading("Dashboard 3 — Mùa vụ và giai đoạn COVID-19", 1)
doc.add_paragraph("Câu hỏi dẫn dắt: các chỉ số thay đổi theo lịch và cùng thời điểm khủng hoảng y tế như thế nào?")
insight(
    8, "Chu kỳ 12 tháng: danceability và acousticness",
    "Gộp các tuần Global theo tháng lịch, danceability gần 0,69–0,70 trong tháng 1–11 nhưng giảm còn 0,661 ở tháng 12. "
    "Acousticness cùng lúc tăng từ khoảng 0,22–0,24 lên 0,294.",
    "Sự quay lại của danh mục nhạc dịp cuối năm là một cơ chế hợp lý cần kiểm chứng bằng tỷ trọng bài hát Giáng sinh hoặc ca khúc cũ trong BXH tháng 12.",
    "Tháng 12 là điểm lệch mùa vụ rõ nhất của hai đặc trưng: âm nhạc trên BXH bớt dễ nhảy và có xu hướng acoustic hơn.",
    "Đây là trung bình gộp các năm có dữ liệu; không thể quy toàn bộ biến động cho Giáng sinh khi chưa kiểm tra từng năm, từng quốc gia và thành phần ca khúc."
)
insight(
    9, "Âm thanh và tử vong COVID-19 theo tuần",
    "Trong 130 tuần Global ghép khớp từ 23/01/2020 đến 14/07/2022, tử vong đạt đỉnh 102.385 ca/tuần ngày 28/01/2021. "
    "Ở tuần này acousticness là 0,283, energy 0,607 và valence 0,513. Trong toàn kỳ, acousticness nằm trong khoảng 0,214–0,367; "
    "energy 0,576–0,668; valence 0,480–0,581.",
    "Âm nhạc nhẹ hơn trong một số giai đoạn có thể phù hợp với thay đổi sinh hoạt tại nhà, hạn chế tụ tập hoặc thay đổi cơ cấu bài hát trên BXH; "
    "dữ liệu hiện không có biến chính sách phong tỏa, hoàn cảnh nghe hoặc tâm trạng người dùng.",
    "Thông điệp thận trọng là “có những thời điểm êm hơn nhưng không nhất thiết buồn hơn”: valence ở đỉnh tử vong vẫn trên 0,50, "
    "song không duy trì trên 0,50 trong toàn kỳ. Hình hai trục cho thấy sự cùng biến động theo thời gian, không tự nó là phép kiểm định tương quan.",
    "Không thể nói tử vong gây ra thay đổi gu nghe hay âm nhạc giúp tự xoa dịu. Cần kiểm soát mùa vụ, xu hướng năm, độ trễ, "
    "thành phần quốc gia và ca khúc; đồ thị hiện chưa thực hiện các kiểm định này."
)

doc.add_heading("Tổng hợp câu chuyện dữ liệu", 1)
doc.add_paragraph(
    "Từ “bản hit có hồ sơ gì?” đến “mỗi thị trường nghe gì?” rồi “hồ sơ ấy đổi ra sao theo thời gian?”, "
    "ba dashboard tạo một tuyến phân tích nhất quán. Nhóm đầu BXH chỉ khác vừa phải về vài chỉ số âm thanh; "
    "nghệ sĩ và hình thức hợp tác liên quan tới quy mô stream; thị trường địa lý và mùa vụ làm thay đổi rõ bối cảnh đọc các chỉ số. "
    "Vì vậy, không nên tìm một công thức hit phổ quát bằng một biến đơn lẻ."
)
doc.add_paragraph(
    "Để chuyển từ kể chuyện khám phá sang kết luận kiểm định hoặc dự báo, bước tiếp theo là mô hình hóa ở cấp ca khúc–tuần: "
    "xác định biến đích (ví dụ vào Top 10 tuần kế tiếp), tách tập huấn luyện/kiểm tra theo thời gian, "
    "kiểm soát nghệ sĩ, quốc gia, mùa vụ và mức độ hiển thị; báo cáo sai số dự báo và tránh rò rỉ dữ liệu từ các tuần tương lai."
)

doc.add_heading("Nguồn và quy tắc diễn giải", 1)
doc.add_paragraph(
    "Nguồn nội bộ: notebooks/04_story.ipynb; data/processed/tracks_dimension.csv, charts_clean.csv, mart_artist_chart.csv, "
    "mart_country_year.csv, mart_audio_trend_weekly.csv và covid_weekly_clean.csv. Các số liệu trong báo cáo được đối chiếu lại trực tiếp "
    "từ các bảng này. “Cơ chế khả dĩ” là giả thuyết diễn giải, không phải kết quả ước lượng tác động nhân quả."
)

doc.core_properties.title = "Khai phá insight từ ba dashboard Spotify"
doc.core_properties.author = "Gia Hân"
doc.core_properties.subject = "Storytelling khoa học về đặc trưng âm thanh và Spotify Charts"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
